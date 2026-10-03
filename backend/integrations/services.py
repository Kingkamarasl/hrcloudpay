import csv
import io
import re
from datetime import datetime
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.utils import timezone

from accounts.audit import audit
from employees.models import Department, Employee


MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 5000
ALLOWED_EXTENSIONS = {'.csv', '.xlsx'}

FIELD_ALIASES = {
    'employee_code': {'employee code', 'employee id', 'employee_id', 'emp id', 'emp no', 'emp no.', 'staff id', 'staff no', 'staff number', 'code'},
    'first_name': {'first name', 'firstname', 'given name', 'givenname'},
    'last_name': {'last name', 'lastname', 'surname', 'family name'},
    'full_name': {'full name', 'name', 'employee name', 'staff name', 'employee'},
    'email': {'email', 'email address', 'work email', 'emailaddress'},
    'phone': {'phone', 'phone number', 'mobile', 'mobile number', 'telephone'},
    'job_title': {'job title', 'title', 'position', 'role', 'job'},
    'department': {'department', 'dept', 'dept.', 'team', 'division'},
    'base_salary': {'salary', 'base salary', 'basic salary', 'monthly pay', 'monthly salary', 'pay', 'wage'},
    'hire_date': {'hire date', 'start date', 'employment date', 'date hired', 'joining date', 'employment date'},
    'employment_status': {'employment status', 'status', 'employee status'},
    'id_card_no': {'id card', 'id card no', 'id card number', 'national id', 'national id no', 'identification number'},
}

TARGET_FIELDS = [
    ('employee_code', 'Employee ID', True),
    ('first_name', 'First name', True),
    ('last_name', 'Last name', True),
    ('email', 'Email', True),
    ('phone', 'Phone', False),
    ('job_title', 'Job title', False),
    ('department', 'Department', False),
    ('base_salary', 'Base salary', True),
    ('hire_date', 'Hire date', False),
    ('employment_status', 'Employment status', False),
    ('id_card_no', 'ID card / identifier', False),
]

STATUS_VALUES = {choice[0] for choice in Employee._meta.get_field('employment_status').choices}


def _normalize_header(value):
    return re.sub(r'[^a-z0-9]+', ' ', str(value or '').strip().lower()).strip()


def suggest_mapping(headers):
    normalized = {_normalize_header(h): h for h in headers}
    mapping = {}
    for field, _label, _required in TARGET_FIELDS:
        aliases = FIELD_ALIASES.get(field, set())
        for alias in aliases:
            if _normalize_header(alias) in normalized:
                mapping[field] = normalized[_normalize_header(alias)]
                break
    return mapping


def _cell_value(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_file(uploaded):
    name = uploaded.name or ''
    lower = name.lower()
    if uploaded.size > MAX_FILE_BYTES:
        raise ValueError('File is too large. Maximum import size is 10 MB.')
    if lower.endswith('.csv'):
        raw = uploaded.read()
        text = raw.decode('utf-8-sig', errors='replace')
        reader = csv.DictReader(io.StringIO(text))
        if not reader.fieldnames:
            raise ValueError('The CSV file does not contain a header row.')
        headers = [h.strip() if h else '' for h in reader.fieldnames]
        rows = []
        for index, row in enumerate(reader, start=2):
            if len(rows) >= MAX_ROWS:
                raise ValueError(f'Import exceeds the {MAX_ROWS:,}-row limit.')
            rows.append({h: _cell_value(row.get(h, '')) for h in headers})
        return headers, rows, raw
    if lower.endswith('.xlsx'):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError('Excel import support is unavailable on this server. Install the openpyxl dependency.') from exc
        uploaded.seek(0)
        raw = uploaded.read()
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        iterator = ws.iter_rows(values_only=True)
        try:
            header_row = next(iterator)
        except StopIteration:
            raise ValueError('The Excel workbook is empty.')
        headers = [_cell_value(value) for value in header_row]
        if not any(headers):
            raise ValueError('The Excel workbook does not contain a header row.')
        rows = []
        for row_values in iterator:
            if len(rows) >= MAX_ROWS:
                raise ValueError(f'Import exceeds the {MAX_ROWS:,}-row limit.')
            values = list(row_values) + [''] * max(0, len(headers) - len(row_values))
            rows.append({headers[i]: _cell_value(values[i]) for i in range(len(headers))})
        wb.close()
        return headers, rows, raw
    raise ValueError('Unsupported file type. Upload a CSV or .xlsx workbook.')


def _parse_date(value):
    if not value:
        return None
    text = str(value).strip()
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%m/%d/%Y', '%d-%m-%Y', '%m-%d-%Y', '%Y/%m/%d'):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def _parse_salary(value):
    if value in ('', None):
        return None
    cleaned = re.sub(r'[^0-9.\-]', '', str(value).replace(',', ''))
    try:
        amount = Decimal(cleaned)
    except (InvalidOperation, ValueError):
        return None
    return amount if amount >= 0 else None


def _split_name(full_name):
    parts = str(full_name or '').strip().split()
    if len(parts) < 2:
        return (parts[0], '') if parts else ('', '')
    return parts[0], ' '.join(parts[1:])


def normalize_row(source, mapping):
    value = lambda field: str(source.get(mapping.get(field, ''), '') or '').strip() if mapping.get(field) else ''
    full_name = value('full_name')
    first_name = value('first_name')
    last_name = value('last_name')
    if full_name and (not first_name or not last_name):
        first_name, last_name = _split_name(full_name)
    salary_raw = value('base_salary')
    salary = _parse_salary(salary_raw)
    hire_raw = value('hire_date')
    hire_date = _parse_date(hire_raw)
    status = value('employment_status').lower().replace(' ', '_') or 'active'
    aliases = {'employed': 'active', 'current': 'active', 'onleave': 'on_leave', 'on leave': 'on_leave', 'left': 'resigned'}
    status = aliases.get(status, status)
    return {
        'employee_code': value('employee_code'),
        'first_name': first_name,
        'last_name': last_name,
        'email': value('email'),
        'phone': value('phone'),
        'job_title': value('job_title'),
        'department': value('department'),
        'base_salary': str(salary) if salary is not None else '',
        'hire_date': hire_date.isoformat() if hire_date else '',
        'employment_status': status,
        'id_card_no': value('id_card_no'),
    }


def validate_rows(company, rows):
    seen_codes = {}
    seen_emails = {}
    output = []
    for number, normalized in enumerate(rows, start=2):
        errors, warnings = [], []
        code = normalized.get('employee_code', '').strip()
        email = normalized.get('email', '').strip()
        if not code:
            errors.append('Employee ID is required.')
        if not normalized.get('first_name'):
            errors.append('First name is required.')
        if not normalized.get('last_name'):
            errors.append('Last name is required.')
        if not email:
            errors.append('Email is required.')
        elif '@' not in email or '.' not in email.split('@')[-1]:
            errors.append('Email address is invalid.')
        if not normalized.get('base_salary'):
            errors.append('Base salary is required and must be numeric.')
        if normalized.get('employment_status') not in STATUS_VALUES:
            errors.append('Employment status is invalid.')
        existing = None
        if code:
            existing = Employee.objects.filter(company=company, employee_code=code).first()
        if existing:
            warnings.append(f'Employee ID already exists ({existing.full_name}).')
        if email:
            email_existing = Employee.objects.filter(company=company, email__iexact=email).first()
            if email_existing and (not existing or email_existing.id != existing.id):
                warnings.append(f'Email already belongs to {email_existing.full_name}.')
                existing = existing or email_existing
        if code and code in seen_codes:
            errors.append(f'Duplicate Employee ID in import (row {seen_codes[code]}).')
        elif code:
            seen_codes[code] = number
        if email and email.lower() in seen_emails:
            errors.append(f'Duplicate email in import (row {seen_emails[email.lower()]}).')
        elif email:
            seen_emails[email.lower()] = number
        if normalized.get('hire_date') == '' and normalized.get('hire_date') is not None:
            warnings.append('Hire date is empty; it can be added later.')
        output.append({
            'row_number': number,
            'data': normalized,
            'errors': errors,
            'warnings': warnings,
            'duplicate_employee_id': existing.id if existing else None,
            'suggested_action': 'update' if existing and not errors else 'create',
        })
    return output


def execute_import(job, request, update_duplicates=False):
    if job.status != 'ready':
        raise ValueError('Only a ready import job can be executed.')
    rows = job.rows
    valid = [r for r in rows if not r.get('errors')]
    company = job.company
    created_ids, updated_ids, department_ids = [], [], []
    skipped = []
    failed = []
    with transaction.atomic():
        job.status = 'importing'
        job.save(update_fields=['status', 'updated_at'])
        current_count = Employee.objects.filter(company=company).count()
        limit = company.employee_limit
        for row in valid:
            data = row['data']
            existing = None
            if data.get('employee_code'):
                existing = Employee.objects.filter(company=company, employee_code=data['employee_code']).first()
            if not existing and data.get('email'):
                existing = Employee.objects.filter(company=company, email__iexact=data['email']).first()
            if existing and not update_duplicates:
                skipped.append(row['row_number'])
                continue
            if not existing and limit is not None and current_count >= limit:
                failed.append({'row_number': row['row_number'], 'reason': f"Your '{company.plan}' plan allows up to {limit} employees."})
                continue
            try:
                dept = None
                if data.get('department'):
                    dept, created_department = Department.objects.get_or_create(company=company, name=data['department'])
                    if created_department and dept.id not in department_ids:
                        department_ids.append(dept.id)
                payload = {
                    'employee_code': data['employee_code'],
                    'first_name': data['first_name'],
                    'last_name': data['last_name'],
                    'email': data['email'],
                    'phone': data.get('phone', ''),
                    'job_title': data.get('job_title', ''),
                    'department': data.get('department', ''),
                    'department_obj': dept,
                    'base_salary': Decimal(data['base_salary']),
                    'hire_date': _parse_date(data.get('hire_date')),
                    'employment_status': data.get('employment_status') or 'active',
                    'id_card_no': data.get('id_card_no', ''),
                }
                if existing:
                    for key, value in payload.items():
                        setattr(existing, key, value)
                    existing.save()
                    updated_ids.append(existing.id)
                else:
                    obj = Employee.objects.create(company=company, **payload)
                    created_ids.append(obj.id)
                    current_count += 1
            except Exception as exc:
                failed.append({'row_number': row['row_number'], 'reason': str(exc)[:240]})
        if failed and not created_ids and not updated_ids and not skipped:
            raise ValueError('No records could be imported.')
        job.status = 'completed_with_warnings' if failed or skipped else 'completed'
        job.result = {
            'created_employee_ids': created_ids,
            'updated_employee_ids': updated_ids,
            'created_department_ids': department_ids,
            'skipped_rows': skipped,
            'failed_rows': failed,
            'update_duplicates': bool(update_duplicates),
        }
        job.completed_at = timezone.now()
        job.save(update_fields=['status', 'result', 'completed_at', 'updated_at'])
    audit(request.user, 'integration_import', f'Imported {len(created_ids)} employees from {job.filename}', company, 'import_job', job.id, {
        'source_type': job.source_type, 'created': len(created_ids), 'updated': len(updated_ids), 'skipped': len(skipped), 'failed': len(failed)
    }, request=request)
    return job


def rollback_import(job, request):
    if job.status not in ('completed', 'completed_with_warnings'):
        raise ValueError('Only a completed import can be rolled back.')
    result = job.result or {}
    created_ids = result.get('created_employee_ids', [])
    department_ids = result.get('created_department_ids', [])
    with transaction.atomic():
        deleted_employees = Employee.objects.filter(company=job.company, id__in=created_ids).delete()[0]
        # Only remove departments that became empty. Never remove pre-existing departments.
        deleted_departments = Department.objects.filter(company=job.company, id__in=department_ids, employees__isnull=True).delete()[0]
        result['rollback'] = {'deleted_employees': deleted_employees, 'deleted_departments': deleted_departments}
        job.result = result
        job.status = 'rolled_back'
        job.rolled_back_at = timezone.now()
        job.save(update_fields=['result', 'status', 'rolled_back_at', 'updated_at'])
    audit(request.user, 'integration_rollback', f'Rolled back import {job.id}', job.company, 'import_job', job.id, result['rollback'], request=request)
    return job
