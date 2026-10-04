"""Spreadsheet import.

The feature exists for a department manager who keeps attendance in a spreadsheet
all month and submits it at the end. That makes it a bulk write into the table
payroll reads, from a file nobody validated, at the end of a period - so the
tests concentrate on the two ways that goes wrong:

- the wrong thing being written silently (a typo'd employee code resolving to
  nobody, and the row vanishing; a blank cell failing to clear a stale time)
- a partial import nobody can account for

The preview/apply split is the safety mechanism, so several tests assert that
preview writes nothing at all.
"""
import csv
import io
from datetime import date, time

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase

from accounts.models import AuditLog
from attendance.importers import TEMPLATE_EXAMPLE
from attendance.models import Attendance

from .tests import AttendanceTestCase

PREVIEW = '/api/attendance/imports/preview/'
APPLY = '/api/attendance/imports/apply/'
TEMPLATE = '/api/attendance/imports/template/'

HEADERS = ['employee_code', 'date', 'status', 'check_in', 'check_out',
           'crossed_midnight', 'notes']


def csv_upload(rows, headers=None, name='attendance.csv'):
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(headers or HEADERS)
    for row in rows:
        writer.writerow(row)
    return SimpleUploadedFile(
        name, buffer.getvalue().encode('utf-8'),
        content_type='text/csv',
    )


def xlsx_upload(rows, headers=None, name='attendance.xlsx'):
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers or HEADERS)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return SimpleUploadedFile(
        name, buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.'
                     'spreadsheetml.sheet',
    )


class ImportTestCase(AttendanceTestCase):
    def setUp(self):
        super().setUp()
        # AttendanceTestCase's two employees are S1 (Sales) and E1 (Engineering).
        self.rows = [
            ['S1', '2026-09-01', 'present', '09:00', '17:00', 'no', ''],
            ['S1', '2026-09-02', 'present', '09:00', '19:30', 'no', 'Long day'],
            ['S1', '2026-09-03', 'absent', '', '', 'no', 'No contact'],
        ]

    def preview(self, upload=None, **payload):
        data = {'file': upload or csv_upload(self.rows)}
        return self.client.post(PREVIEW, data, format='multipart', **payload)

    def apply(self, token, **extra):
        return self.client.post(APPLY, {'preview_token': token, **extra},
                                format='json')


# ---------------------------------------------------------------------------
# Preview writes nothing
# ---------------------------------------------------------------------------
class PreviewIsSafeTests(ImportTestCase):
    def test_preview_creates_no_records(self):
        response = self.preview()

        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(Attendance.objects.count(), 0)

    def test_preview_creates_no_audit_row(self):
        self.preview()
        self.assertEqual(AuditLog.objects.filter(target_type='attendance').count(), 0)

    def test_preview_reports_what_it_would_do(self):
        body = self.preview().json()

        self.assertEqual(body['summary'], {'create': 3, 'update': 0,
                                          'unchanged': 0, 'error': 0})
        self.assertEqual(body['total_rows'], 3)
        self.assertTrue(body['preview_token'])


# ---------------------------------------------------------------------------
# Formats and headings
# ---------------------------------------------------------------------------
class SpreadsheetFormatTests(ImportTestCase):
    def test_an_xlsx_upload_is_read(self):
        body = self.preview(xlsx_upload(self.rows)).json()
        self.assertEqual(body['summary']['create'], 3)

    def test_date_cells_that_arrive_as_real_dates_are_read(self):
        rows = [['S1', date(2026, 9, 1), 'present', time(9, 0), time(17, 0),
                 False, '']]
        body = self.preview(xlsx_upload(rows)).json()

        self.assertEqual(body['summary']['create'], 1)
        self.assertEqual(body['rows'][0]['date'], '2026-09-01')

    def test_common_header_spellings_are_recognised(self):
        headers = ['Employee Code', 'Date', 'Status', 'Clock In', 'Clock Out',
                   'Next Day', 'Remarks']
        body = self.preview(xlsx_upload(self.rows, headers=headers)).json()
        self.assertEqual(body['summary']['create'], 3)

    def test_a_missing_required_column_is_refused_with_guidance(self):
        headers = ['date', 'status']
        response = self.preview(xlsx_upload(self.rows, headers=headers))

        self.assertEqual(response.status_code, 400)
        self.assertIn('employee', str(response.data).lower())

    def test_an_unsupported_file_type_is_refused(self):
        upload = SimpleUploadedFile('attendance.pdf', b'%PDF-1.4',
                                    content_type='application/pdf')
        response = self.preview(upload)

        self.assertEqual(response.status_code, 400)
        self.assertIn('xlsx', str(response.data).lower())

    def test_an_empty_file_is_refused(self):
        response = self.preview(csv_upload([]))
        self.assertEqual(response.status_code, 400)

    def test_a_missing_file_is_refused(self):
        response = self.client.post(PREVIEW, {}, format='multipart')
        self.assertEqual(response.status_code, 400)

    def test_the_template_is_a_downloadable_xlsx(self):
        response = self.client.get(TEMPLATE)

        self.assertEqual(response.status_code, 200)
        self.assertIn('spreadsheetml', response['Content-Type'])

        from openpyxl import load_workbook
        sheet = load_workbook(io.BytesIO(response.content)).worksheets[0]
        header = [c.value for c in sheet[1]]
        self.assertEqual(header, HEADERS)

    def test_the_template_ships_examples_on_a_separate_sheet(self):
        """An untouched template must not look like a broken file.

        With examples on the sheet that gets filled in, uploading the template
        unchanged previews as rows that all fail to resolve to an employee -
        which teaches the manager the format is broken by showing them its own
        error output.
        """
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(self.client.get(TEMPLATE).content))
        self.assertEqual([c.value for c in workbook.worksheets[0][1]], HEADERS)
        self.assertEqual(workbook.worksheets[0].max_row, 1)
        self.assertEqual(workbook.worksheets[1].title, 'Example')

    def test_uploading_the_untouched_template_says_so_plainly(self):
        upload = SimpleUploadedFile('attendance.xlsx',
                                    self.client.get(TEMPLATE).content)
        response = self.preview(upload)

        self.assertEqual(response.status_code, 400)
        self.assertIn('no data rows', str(response.data).lower())

    def test_the_example_rows_parse_once_the_employees_exist(self):
        """The examples must be valid input, or they demonstrate the wrong thing.

        The base fixture already owns S1 and E1, so the examples are remapped
        onto those rather than creating codes that already exist.
        """
        self.make_employee(self.acme, 'EMP-003', 'Night', 'Shift')
        remap = {'EMP-001': 'S1', 'EMP-002': 'E1'}
        rows = [[remap.get(row[0], row[0]), *row[1:]] for row in TEMPLATE_EXAMPLE]

        body = self.preview(xlsx_upload(rows)).json()

        self.assertEqual(body['summary']['error'], 0, body['rows'])
        self.assertEqual(body['summary']['create'], len(TEMPLATE_EXAMPLE))
        # The night-shift example has to survive as a real overnight shift.
        night = [r for r in body['rows'] if r['crossed_midnight']]
        self.assertEqual(len(night), 1)


# ---------------------------------------------------------------------------
# Row validation - never fail on the first bad line
# ---------------------------------------------------------------------------
class RowValidationTests(ImportTestCase):
    def test_an_unknown_employee_code_is_reported_not_skipped_silently(self):
        rows = list(self.rows) + [['NOPE', '2026-09-04', 'present', '', '', 'no', '']]
        body = self.preview(csv_upload(rows)).json()

        bad = [r for r in body['rows'] if r['action'] == 'error']
        self.assertEqual(len(bad), 1)
        self.assertIn('NOPE', bad[0]['errors'][0])
        # Still planned as creates for the good rows, so the manager sees the
        # full picture rather than a short count.
        self.assertEqual(body['summary']['create'], 3)

    def test_every_bad_row_is_reported_not_just_the_first(self):
        rows = [
            ['NOPE', '2026-09-04', 'present', '', '', 'no', ''],
            ['ALSOMISSING', '2026-09-05', 'present', '', '', 'no', ''],
        ]
        body = self.preview(csv_upload(rows)).json()
        self.assertEqual(body['summary']['error'], 2)

    def test_an_unparseable_date_is_reported_with_its_row_number(self):
        rows = [['S1', 'the seventh', 'present', '', '', 'no', '']]
        body = self.preview(csv_upload(rows)).json()

        bad = body['rows'][0]
        self.assertEqual(bad['row'], 2)
        self.assertIn('date', bad['errors'][0])

    def test_an_unknown_status_is_reported_rather_than_guessed(self):
        """Guessing here means the wrong pay, so it has to be a row error."""
        rows = [['S1', '2026-09-01', 'on holiday', '', '', 'no', '']]
        body = self.preview(csv_upload(rows)).json()

        self.assertEqual(body['summary']['error'], 1)
        self.assertIn('status', body['rows'][0]['errors'][0])

    def test_status_synonyms_are_accepted(self):
        rows = [['S1', '2026-09-01', 'Present', '', '', 'no', '']]
        body = self.preview(csv_upload(rows)).json()
        self.assertEqual(body['summary']['create'], 1)

    def test_reversed_times_are_reported(self):
        rows = [['S1', '2026-09-01', 'present', '17:00', '09:00', 'no', '']]
        body = self.preview(csv_upload(rows)).json()

        self.assertEqual(body['summary']['error'], 1)
        self.assertIn('earlier than the check-in', body['rows'][0]['errors'][0])

    def test_an_overnight_shift_is_accepted_with_the_flag(self):
        rows = [['S1', '2026-09-05', 'present', '22:00', '06:00', 'yes', '']]
        body = self.preview(csv_upload(rows)).json()

        self.assertEqual(body['summary']['create'], 1, body['rows'])
        self.assertTrue(body['rows'][0]['crossed_midnight'])

    def test_the_flag_contradicting_the_times_is_reported(self):
        rows = [['S1', '2026-09-01', 'present', '09:00', '17:00', 'yes', '']]
        body = self.preview(csv_upload(rows)).json()
        self.assertEqual(body['summary']['error'], 1)

    def test_a_duplicate_employee_and_date_in_one_file_is_reported(self):
        rows = list(self.rows) + [['S1', '2026-09-01', 'absent', '', '', 'no', '']]
        body = self.preview(csv_upload(rows)).json()

        self.assertEqual(body['summary']['error'], 1)
        self.assertIn('more than once', body['rows'][3]['errors'][0])

    def test_a_blank_status_assumes_present_and_says_so(self):
        rows = [['S1', '2026-09-01', '', '', '', 'no', '']]
        body = self.preview(csv_upload(rows)).json()

        self.assertEqual(body['rows'][0]['status'], 'present')
        self.assertIn('assuming present', body['rows'][0]['errors'][0])


# ---------------------------------------------------------------------------
# Scope
# ---------------------------------------------------------------------------
class ImportScopeTests(ImportTestCase):
    def test_employees_cannot_import(self):
        self.authenticate(self.worker)
        response = self.preview()

        self.assertEqual(response.status_code, 403)
        self.assertEqual(Attendance.objects.count(), 0)

    def test_a_department_manager_imports_their_own_department(self):
        manager = self.make_user(self.acme, 'sales-mgr', 'department_manager',
                                 managed_department='Sales')
        self.authenticate(manager)

        body = self.preview().json()
        self.assertEqual(body['summary']['create'], 3)

    def test_a_department_manager_cannot_import_another_department(self):
        manager = self.make_user(self.acme, 'sales-mgr', 'department_manager',
                                 managed_department='Sales')
        self.authenticate(manager)

        rows = [['E1', '2026-09-01', 'present', '09:00', '17:00', 'no', '']]
        body = self.preview(csv_upload(rows)).json()

        self.assertEqual(body['summary']['create'], 0)
        self.assertEqual(body['summary']['error'], 1)
        self.assertIn('cannot record attendance', body['rows'][0]['errors'][0])

    def test_another_tenants_employee_code_does_not_resolve(self):
        """A code from another company must be "no match", not silently ours."""
        rows = [['X1', '2026-09-01', 'present', '09:00', '17:00', 'no', '']]
        body = self.preview(csv_upload(rows)).json()

        self.assertEqual(body['summary']['create'], 0)
        self.assertEqual(body['rows'][0]['employee_id'], None)

    def test_a_manager_who_cannot_apply_what_they_previewed_is_refused(self):
        """Reach is re-checked at apply, not trusted from preview time."""
        manager = self.make_user(self.acme, 'sales-mgr', 'department_manager',
                                 managed_department='Sales')
        self.authenticate(manager)
        token = self.preview().json()['preview_token']

        # Reach changes between the two calls.
        manager.managed_department = 'Engineering'
        manager.save()

        response = self.apply(token)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['created'], 0)
        self.assertEqual(response.data['refused'], 3)
        self.assertEqual(Attendance.objects.count(), 0)


# ---------------------------------------------------------------------------
# Apply
# ---------------------------------------------------------------------------
class ApplyTests(ImportTestCase):
    def test_apply_creates_the_planned_records(self):
        token = self.preview().json()['preview_token']

        response = self.apply(token)

        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['created'], 3)
        self.assertEqual(Attendance.objects.count(), 3)

    def test_apply_stores_the_values_the_preview_showed(self):
        token = self.preview().json()['preview_token']
        self.apply(token)

        long_day = Attendance.objects.get(date=date(2026, 9, 2))
        self.assertEqual(long_day.status, 'present')
        self.assertEqual(long_day.check_in, time(9, 0))
        self.assertEqual(long_day.check_out, time(19, 30))

    def test_apply_is_audited(self):
        token = self.preview().json()['preview_token']
        self.apply(token)

        entry = AuditLog.objects.filter(
            action='create', target_type='attendance').order_by('-id').first()
        self.assertIsNotNone(entry)
        self.assertIn('Imported attendance from', entry.message)
        self.assertEqual(entry.metadata['created'], 3)

    def test_a_token_cannot_be_applied_twice(self):
        """Replaying one upload would double every record in it."""
        token = self.preview().json()['preview_token']
        self.apply(token)

        second = self.apply(token)

        self.assertEqual(second.status_code, 400)
        self.assertEqual(Attendance.objects.count(), 3)

    def test_an_unknown_token_is_refused(self):
        response = self.apply('not-a-real-token')
        self.assertEqual(response.status_code, 400)

    def test_apply_without_a_token_is_refused(self):
        response = self.client.post(APPLY, {}, format='json')
        self.assertEqual(response.status_code, 400)

    def test_another_users_token_cannot_be_applied(self):
        """The token is namespaced per user, so a leaked one is useless."""
        token = self.preview().json()['preview_token']

        other = self.make_user(self.acme, 'other-hr', 'hr')
        self.authenticate(other)
        response = self.apply(token)

        self.assertEqual(response.status_code, 400)
        self.assertEqual(Attendance.objects.count(), 0)

    def test_a_resubmitted_identical_file_changes_nothing(self):
        token = self.preview().json()['preview_token']
        self.apply(token)

        again = self.preview()
        body = again.json()
        self.assertEqual(body['summary']['unchanged'], 3)
        self.assertEqual(body['summary']['create'], 0)

        applied = self.apply(body['preview_token'])
        self.assertEqual(applied.data['created'], 0)
        self.assertEqual(applied.data['updated'], 0)
        self.assertEqual(Attendance.objects.count(), 3)

    def test_a_corrected_resubmission_updates_rather_than_duplicating(self):
        token = self.preview().json()['preview_token']
        self.apply(token)

        corrected = [['S1', '2026-09-03', 'present', '09:00', '17:00', 'no', '']]
        body = self.preview(csv_upload(corrected)).json()
        self.assertEqual(body['summary']['update'], 1)

        applied = self.apply(body['preview_token']).json()

        self.assertEqual(applied['updated'], 1)
        self.assertEqual(applied['created'], 0)
        self.assertEqual(Attendance.objects.count(), 3)
        record = Attendance.objects.get(date=date(2026, 9, 3))
        self.assertEqual(record.status, 'present')

    def test_a_blank_cell_clears_a_time_that_was_there(self):
        """Preview and apply must not disagree about what a blank means.

        If apply skipped None values, a sheet with an empty check-in would leave
        the old time in place while the preview said it would change.
        """
        self.record(self.sales, date='2026-09-01', check_in=time(9, 0),
                    check_out=time(17, 0))

        blanked = [['S1', '2026-09-01', 'present', '', '', 'no', '']]
        body = self.preview(csv_upload(blanked)).json()
        self.assertEqual(body['summary']['update'], 1)

        self.apply(body['preview_token'])

        record = Attendance.objects.get(employee=self.sales, date=date(2026, 9, 1))
        self.assertIsNone(record.check_in)
        self.assertIsNone(record.check_out)

    def test_error_rows_are_refused_not_partially_applied(self):
        rows = list(self.rows) + [['NOPE', '2026-09-04', 'present', '', '', 'no', '']]
        token = self.preview(csv_upload(rows)).json()['preview_token']

        response = self.apply(token).json()

        self.assertEqual(response['created'], 3)
        self.assertEqual(response['refused'], 1)
        self.assertEqual(response['refused_rows'][0]['employee_ref'], 'NOPE')

    def test_an_overnight_shift_survives_the_round_trip(self):
        rows = [['S1', '2026-09-05', 'present', '22:00', '06:00', 'yes', 'Night']]
        token = self.preview(csv_upload(rows)).json()['preview_token']

        self.apply(token)

        record = Attendance.objects.get()
        self.assertTrue(record.crossed_midnight)
        self.assertEqual(record.worked_minutes, 8 * 60)

    def test_an_oversized_file_is_refused_before_it_is_read(self):
        upload = SimpleUploadedFile('attendance.csv', b'x' * (5 * 1024 * 1024 + 10),
                                    content_type='text/csv')
        response = self.preview(upload)
        # 413 from Django's upload handling, 400 from the view's own check -
        # either is a refusal, and neither writes anything.
        self.assertIn(response.status_code, (400, 413))
        self.assertEqual(Attendance.objects.count(), 0)