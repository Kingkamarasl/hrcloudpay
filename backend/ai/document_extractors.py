from pathlib import Path

ALLOWED_EXTENSIONS = {'.pdf', '.docx'}
MAX_FILE_SIZE = 10 * 1024 * 1024

# pytesseract is a wrapper around the `tesseract` binary, not an implementation
# of OCR. Installing the Python package on a host with no such binary achieves
# nothing, and shared hosting gives no root and no package manager. Worth saying
# in one place, because "the dependency is installed" and "OCR works" are very
# different claims.
OCR_UNAVAILABLE = 'no-ocr-engine'


def _open_for_ocr(uploaded_file):
    """Open the PDF with PyMuPDF, for rasterising pages.

    Returns (document, reason). `reason` is None on success. Opened lazily
    because a PDF whose pages all carry a text layer never needs it, and a
    host without PyMuPDF should still be able to read digital documents.
    """
    try:
        import fitz
    except ImportError:
        return None, (
            'This server cannot read scanned pages: the PyMuPDF library is not '
            'installed.'
        )

    try:
        uploaded_file.seek(0)
        return fitz.open(stream=uploaded_file.read(), filetype='pdf'), None
    except Exception as exc:
        return None, 'The PDF could not be opened for scanning (%s).' % type(exc).__name__


def _ocr_page(pdf, index):
    """OCR one page. Returns (text, reason); reason is None on success.

    The two failure kinds are kept apart on purpose. "This server has no OCR
    engine" is something an administrator fixes once, for everyone. "OCR ran and
    still could not read this page" is something the uploader fixes by
    re-scanning, and it happens on perfectly healthy servers - a skewed
    page, a low-resolution photograph, handwriting. Reporting them identically
    would send people to the wrong place.
    """
    try:
        import pytesseract
    except ImportError:
        return '', OCR_UNAVAILABLE

    try:
        import fitz
        from PIL import Image
        import io

        pix = pdf[index].get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
        image = Image.open(io.BytesIO(pix.tobytes('png')))
        return (pytesseract.image_to_string(image) or '').strip(), None
    except pytesseract.TesseractNotFoundError:
        # The Python package is installed; the binary it drives is not.
        return '', OCR_UNAVAILABLE
    except Exception as exc:
        return '', 'OCR could not read the page (%s).' % type(exc).__name__


def _format_pages(numbers):
    """[1, 2, 3, 7] -> '1-3, 7'. A 40-page scan with 40 unreadable pages should
    not produce a 200-character error message."""
    numbers = sorted(numbers)
    if not numbers:
        return ''

    runs = []
    start = previous = numbers[0]
    for value in numbers[1:]:
        if value == previous + 1:
            previous = value
            continue
        runs.append((start, previous))
        start = previous = value
    runs.append((start, previous))

    return ', '.join(
        str(first) if first == last else '%d-%d' % (first, last)
        for first, last in runs
    )


def _unreadable_message(meta):
    listed = _format_pages(meta['unreadable_pages'])
    total = meta['page_count']
    count = len(meta['unreadable_pages'])
    pages = 'Page' if count == 1 else 'Pages'

    reason = meta.get('ocr_unavailable_reason')
    if reason == OCR_UNAVAILABLE:
        cause = ('This server has no OCR engine installed, so pages that are '
                 'images of text rather than text cannot be read.')
    elif reason:
        cause = reason
    else:
        cause = 'Those pages could not be read.'

    return (
        '%s %s of %s could not be read. %s Nothing has been saved - a document '
        'with pages missing would be indexed as though it were whole, and every '
        'answer drawn from it would then be wrong about what it says. Re-save '
        'the file with a text layer (print to PDF rather than scan it), or split '
        'the unreadable pages into a separate document and upload that.'
        % (pages, listed, total, cause)
    )


def _nothing_readable_message(meta):
    """Every page came back empty.

    Naming the reason is the whole point. "No readable text" on its own leaves
    the reader with nothing to act on, and the two causes need opposite
    responses: a missing OCR engine is an administrator's problem and no amount
    of re-scanning will fix it, while a document that extracted badly is
    something the uploader can fix by exporting it properly.
    """
    reason = meta.get('ocr_unavailable_reason')

    if reason == OCR_UNAVAILABLE:
        detail = (' Every page appears to be an image of text, and this server '
                  'has no OCR engine installed, so it cannot be read.')
    elif reason:
        detail = ' %s' % reason
    else:
        detail = ' Every page is blank, or carries no text.'

    return ('No readable text could be extracted from this PDF.%s Nothing has '
            'been saved.' % detail)


def extract_pdf(uploaded_file):
    """Read a PDF page by page, classifying each one.

    A page ends up in exactly one of four states:

    ``text``        read from its embedded text layer
    ``ocr``         read by rasterising it and running OCR
    ``blank``       no text and no images - genuinely empty, nothing lost
    ``unreadable``  images but no text, and OCR could not recover it

    The fourth is the one that used to be invisible. It collapsed to an empty
    string, the caller joined the pages that did work, and the document was
    saved and indexed as though complete. Nothing reported the gap; the AI went
    on to answer questions about content that was simply absent.

    ``blank`` is separated from ``unreadable`` deliberately. The back of a sheet
    is blank in contracts and reports constantly, and treating every empty page
    as a failure would make the feature unusable for ordinary documents.
    """
    from pypdf import PdfReader

    uploaded_file.seek(0)
    reader = PdfReader(uploaded_file)

    pdf = None
    pdf_reason = None
    tried_ocr = False
    pages = []
    unreadable = []
    ocr_used = False
    ocr_reason = None

    try:
        for number, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or '').strip()

            if not text:
                if not tried_ocr:
                    tried_ocr = True
                    pdf, pdf_reason = _open_for_ocr(uploaded_file)

                if pdf is None:
                    # Cannot tell a blank page from a scan, so it cannot be
                    # dismissed as blank. Recorded as unreadable instead.
                    unreadable.append(number)
                    ocr_reason = ocr_reason or pdf_reason
                    pages.append({
                        'page_number': number, 'text': '', 'source': 'unreadable',
                    })
                    continue

                has_images = bool(pdf[number - 1].get_images(full=True))
                if not has_images:
                    pages.append({
                        'page_number': number, 'text': '', 'source': 'blank',
                    })
                    continue

                text, reason = _ocr_page(pdf, number - 1)
                if text:
                    ocr_used = True
                    source = 'ocr'
                else:
                    unreadable.append(number)
                    ocr_reason = ocr_reason or reason
                    source = 'unreadable'
            else:
                source = 'text'

            pages.append({'page_number': number, 'text': text, 'source': source})

        return pages, {
            'ocr_used': ocr_used,
            'page_count': len(pages),
            'unreadable_pages': unreadable,
            'ocr_unavailable_reason': ocr_reason,
        }
    finally:
        if pdf is not None:
            pdf.close()


def _docx_blocks(uploaded_file):
    from docx import Document
    uploaded_file.seek(0)
    doc = Document(uploaded_file)
    blocks = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            style = getattr(paragraph.style, 'name', '') if paragraph.style else ''
            section = text if style and 'heading' in style.lower() else ''
            blocks.append({'text': text, 'section': section})
    for table_index, table in enumerate(doc.tables, start=1):
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append({'text': ' | '.join(cells), 'section': f'Table {table_index}'})
    return blocks


def extract_text(uploaded_file):
    """Extract document text with page/section metadata and OCR fallback.

    Raises ValueError when any part of the document could not be read. Returning
    what was recovered instead would leave the caller to save a document with
    pages missing, and the damage would surface much later as a confident wrong
    answer rather than as an error.
    """
    name = uploaded_file.name or ''
    ext = Path(name).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError('Only PDF and DOCX files are supported.')
    if uploaded_file.size and uploaded_file.size > MAX_FILE_SIZE:
        raise ValueError('The maximum document size is 10 MB.')

    if ext == '.pdf':
        pages, meta = extract_pdf(uploaded_file)
        if not any(page['text'] for page in pages):
            raise ValueError(_nothing_readable_message(meta))
        if meta['unreadable_pages']:
            raise ValueError(_unreadable_message(meta))
        return pages, ext, meta

    blocks = _docx_blocks(uploaded_file)
    if not blocks:
        raise ValueError('No readable text could be extracted from this DOCX document.')
    return blocks, ext, {
        'ocr_used': False,
        'page_count': None,
        'unreadable_pages': [],
        'ocr_unavailable_reason': None,
    }